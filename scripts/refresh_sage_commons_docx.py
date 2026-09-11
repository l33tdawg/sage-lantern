#!/usr/bin/env python3
"""Refresh legacy SAGE Commons Word artifacts without rebuilding their design."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def replace_exact(doc, replacements: dict[str, str]) -> None:
    remaining = dict(replacements)
    for paragraph in iter_paragraphs(doc):
        old = paragraph.text
        if old not in remaining:
            continue
        new = remaining.pop(old)
        if paragraph.runs:
            paragraph.runs[0].text = new
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(new)
    if remaining:
        missing = "\n".join(f"- {text}" for text in remaining)
        raise RuntimeError(f"Expected source text not found:\n{missing}")


COMMON = {
    "Working draft • 28 July 2026 • Architecture revision verified against SAGE v11 documentation":
        "Programme baseline • 11 August 2026 • Aligned with the SAGE Commons master technical PRD",
}


PRD = {
    **COMMON,
    "Sovereign personal memory plus a governed public SAGE chain for disconnected communities":
        "Private local AI plus a governed public knowledge network for disconnected communities",
    "SAGE Commons is an open, nonprofit-oriented platform for communities with limited, intermittent or unaffordable connectivity. Every SAGE Lantern is a complete local AI appliance with local inference, speech-to-text, text-to-speech, a readable colour display, local reference collections and sovereign personal memory.":
        "SAGE Commons is the open mission, governed public knowledge network and community programme for places with limited, intermittent or unaffordable connectivity. SAGE Lantern is its portable local-first AI product: local inference, speech, display, reference collections and sovereign personal memory, plus bounded Commons proposer, query and cache functions.",
    "The architecture uses two distinct logical ledgers. A personal SAGE chain remains device-local and contains private memory. A separate public Commons SAGE chain contains only approved public domains. Knowledge Points run its validator and full-replica roles over normal internet where infrastructure exists. Handhelds and mobile couriers carry proposals, committed public state and approved bundles between encounters; their cached public height may lag without changing the personal chain.":
        "A personal SAGE remains device-local and private. The Lantern does not run a second public SAGE process: an isolated Commons adapter holds a proposal outbox, query client and replaceable public cache. Stable Lightbearer services support the public network; only independently governed Lightbearer Validators hold consensus power. SAGE Lamplighters carry verified packs between connected and disconnected locations.",
    "This separation follows the SAGE implementation rather than redefining federation. In current SAGE, federation connects independent chains for live Read or opt-in Copy, and a copied memory is re-admitted through the receiver's own consensus. Same-chain peer join and state sync share one ledger. SAGE Commons therefore treats the public Commons chain as a distinct network role, not as a magical merge of sovereign personal chains [10–12].":
        "This separation respects upstream SAGE rather than redefining federation. SAGE v11.18.4 federation links independent chains for authenticated Read or opt-in Copy; it is not Commons ingress. Public proposal intake, external query, deterministic frozen-height export and snapshot-pack production remain explicit Commons or required-SAGE-v12 contracts, never capabilities inferred from generic state sync [10–12].",
    "DEVICE-LOCAL PERSONAL SAGE CHAIN\n└── Private domains\n│   ├── user profile\n│   ├── personal conversations\n│   ├── health history\n│   ├── preferences\n│   └── device-specific context\n\nSEPARATE PUBLIC COMMONS SAGE CHAIN (cached locally)\n└── Public domains\n    ├── agriculture\n    ├── medicine\n    ├── education\n    ├── local services\n    └── shared community knowledge":
        "DEVICE-LOCAL PERSONAL SAGE\n└── Private domains\n│   ├── user profile\n│   ├── conversations\n│   ├── health history\n│   ├── preferences\n│   └── device context\n\nISOLATED COMMONS ADAPTER + CACHE\n├── PUBLIC-CANDIDATE outbox\n├── authenticated query client\n└── replaceable public cache\n    ├── agriculture\n    ├── medicine\n    ├── education\n    └── local services",
    "Private and public are policy classes backed by separate chain roles, not merely folders in one replicated ledger. Every domain carries rules for retention, encryption, retrieval, export, access, challenge rights, expiry and revocation. A device may additionally maintain quarantine and transient-working domains, but those do not weaken the private/public boundary.":
        "Private and public are enforced data planes, not folders in one replicated ledger. Personal data is private by default. Sharing creates a separate, editable and sanitized PUBLIC-CANDIDATE; the private source is never reclassified or copied into public canonical state. Policy defines consent, retention, encryption, retrieval, publication, challenge and correction behavior.",
    "Each device has a sovereign personal SAGE role and a separate Commons public-chain role. The personal chain is never joined to the Commons chain. A device may prepare a public memory submission while offline, but it becomes proposed only when broadcast to the Commons chain; after block inclusion, the SAGE validator-vote quorum determines whether it becomes committed public memory [10–12].":
        "Each Lantern has sovereign personal SAGE plus a separately stored Commons adapter, outbox and cache. Explicit owner consent creates a signed logical proposal envelope. Intake verifies and quarantines it; after approval the Lantern reconnects and makes one fresh nonce/timestamp-bound SAGE submission. Only stable Lightbearer Validators determine canonical outcome [10–12].",
    "DEVICE A              DEVICE B              DEVICE C\nPersonal SAGE A        Personal SAGE B        Personal SAGE C\nprivate; local         private; local         private; local\n\n        └── public submissions + committed Commons state ──┐\n                                                           ▼\nKnowledge Point validators/full replicas ← internet → peer Knowledge Points\n                  ONE PUBLIC COMMONS SAGE CHAIN":
        "LANTERN A            LANTERN B            LANTERN C\nPersonal SAGE A        Personal SAGE B        Personal SAGE C\nprivate; local         private; local         private; local\n\n       └── proposals + queries; verified cache/packs ───────┐\n                                                           ▼\nIntake / Lightbearer services → independent Lightbearer Validators\n                    PUBLIC SAGE COMMONS",
    "Private domains are absent from Commons inventories, transaction ranges and snapshots.":
        "Private domains are absent from proposal envelopes, Commons queries, public caches, packs and operational inventories.",
    "Knowledge Points run the public chain's validators and full replicas; validator membership and governance changes remain on-chain SAGE decisions.":
        "Knowledge Points are physical or institutional facilities. They may host isolated Lightbearer profiles, but only separately admitted, independently governed Lightbearer Validators receive consensus power.",
    "A Foundation publisher or field device may submit a candidate public memory, but cannot unilaterally commit it.":
        "A Lantern or authorized publisher may propose public knowledge, but neither the proposer nor the Foundation can unilaterally commit it.",
    "Travelling doctors, teachers, technicians or other couriers may carry a recent public-chain replica/cache and seed local peers; validator authority remains credential-bound.":
        "A SAGE Lamplighter carries verified knowledge, model and runtime packs and seeds local peers; the carrier is untrusted and never gains publishing, Root or validator authority.",
    "LoRa/Wi-Fi adaptation for disconnected proposal carriage and public-chain catch-up is new SAGE Commons work, not a feature already claimed by upstream SAGE.":
        "Disconnected proposal carriage, signed pack distribution, public-cache refresh and LoRa/Wi-Fi coordination are Commons work, not capabilities claimed from upstream SAGE.",
    "Device-local personal SAGE plus a separately isolated Commons public-chain client/cache.":
        "Device-local personal SAGE plus an isolated Commons proposer/query adapter and replaceable public cache.",
    "Signed SAGE transactions, per-node identity, BFT validation, memory votes, challenges, revocation and audit trail.":
        "Signed logical proposal envelopes, fresh SAGE submit authorization, per-node identity, validation, challenges, revocation and receipts.",
    "LoRa announcements, compact public submissions, query requests and public-chain checkpoint summaries.":
        "LoRa availability/control hints, compact research messages and pack/checkpoint summaries within strict airtime limits.",
    "Automatic Wi-Fi rendezvous for transaction ranges, snapshots, evidence attachments and large knowledge/software bundles.":
        "Automatic Wi-Fi rendezvous for verified cache chunks, evidence and signed knowledge, model and runtime packs.",
    "Knowledge Point validator/full-replica services for research, publication, public-chain consensus, internet replication, caching, diagnostics and updates.":
        "Knowledge Point facilities and isolated Lightbearer profiles for research, query/export, intake, caching, diagnostics or independently governed validation.",
    "The local assistant searches the personal SAGE chain and the locally cached committed Commons knowledge as distinct sources.":
        "The local assistant searches personal SAGE and the replaceable, provenance-bearing Commons cache as distinct sources.",
    "If evidence is adequate, the device answers locally; a private memory may enter the personal chain while an eligible public observation is queued as a Commons submission.":
        "If evidence is adequate, the device answers locally. Private memory may enter personal SAGE; an eligible observation becomes public only through explicit consent and a separate PUBLIC-CANDIDATE draft.",
    "If evidence is inadequate, the device creates a minimised public research request for the mesh or a Knowledge Point.":
        "If evidence is inadequate, the device may create a minimised public query for the mesh or a Knowledge Point-hosted Lightbearer service.",
    "A returning result becomes a candidate Commons memory submission, passes block consensus and the SAGE memory-vote quorum, and may become committed in the relevant public domain.":
        "A result remains cache material unless the user separately chooses to publish. Approved proposals pass intake, one fresh Lantern-authored submit and public consensus before becoming committed.",
    "SAGE Lantern: hold the side button, speak, release, and receive an immediate local response. The display shows the transcript, answer, supporting sources, confidence, domain destination and any safety warning. Before a memory is shared, the interface must explain whether it is private, public, pending validation or quarantined.":
        "SAGE Lantern: hold the side button, speak, release and receive a local response. The display shows transcript, sources, confidence, freshness and warnings. Sharing is a separate owner-confirmed flow that clearly labels MY PRIVATE, PUBLIC-CANDIDATE, queued, received, approved, submitted, committed, rejected and stale cache states.",
    "Which Knowledge Points belong to the initial Commons validator set, and how are validator changes governed?":
        "Which independent operators qualify for the initial Lightbearer Validator set, and how are admission and changes governed?",
    "Each SAGE Lantern keeps a sovereign personal SAGE chain for private memory and an intermittent cache/client role on a separate public Commons SAGE chain. Knowledge Point validators govern the public chain; LoRa coordinates and Wi-Fi carries proposals, committed public state and knowledge bundles.":
        "SAGE Commons is the mission and public network. Each Lantern keeps sovereign personal SAGE and only bounded proposer, query and cache functions. Independent Lightbearer Validators govern canonical public state; Lamplighters and local peers distribute verified packs over bulk links while LoRa coordinates.",
    "Mobile Knowledge Point": "SAGE Lamplighter",
    "A travelling doctor, teacher or technician carries a recent validated public chain and approved bundles, then seeds local peers over Wi-Fi.":
        "An untrusted field kit carries verified signed knowledge, model and runtime packs, then seeds local peers over Wi-Fi, Ethernet or USB.",
    "Offers a minimised mesh query or queue for the next Knowledge Point.":
        "Offers a minimised query or queues it for a suitable Lightbearer service.",
    "Use private relevant memory locally without exposing it to peers or Knowledge Points.":
        "Use relevant private memory locally without exposing it to peers, operators, Knowledge Points or Lightbearers.",
    "FR-6 Commons public-chain role": "FR-6 Commons edge role",
    "Queue eligible submissions and exchange committed public-chain state without joining or copying the personal chain.":
        "Queue PUBLIC-CANDIDATE proposals, query authorized services and ingest only verified replaceable cache data.",
    "Exchange missing transaction ranges, snapshots, evidence and large bundles with resume and deduplication.":
        "Exchange verified content-addressed cache chunks, evidence and signed packs with resume and deduplication.",
    "FR-10 Knowledge Points": "FR-10 Knowledge Points / Lightbearers",
    "Operate Commons-chain validators/full replicas for consensus, internet replication, research, publishing and bundle caching.":
        "Host isolated research or Lightbearer services; validators require separate governance, identities, keys and qualification.",
    "Private-domain data does not appear in beacons, inventories, public snapshots or upstream requests.":
        "Private-domain data does not appear in beacons, inventories, public caches/packs, proposals or upstream requests.",
    "Zero private-domain objects observed in mesh traffic, public snapshots or Knowledge Point requests.":
        "Zero private-domain objects observed in mesh traffic, public caches/packs, proposals or service requests.",
    "After Knowledge Point contact, every reachable cache/peer reaches the expected Commons-chain checkpoint while personal-chain heights remain unchanged.":
        "After Lightbearer, Lamplighter or peer contact, each reachable Lantern verifies and activates the expected cache/pack sequence; personal SAGE remains unchanged.",
    "Six devices plus one Knowledge Point; submission/consensus flow; Wi-Fi catch-up; red-team privacy tests.":
        "Six Lanterns plus service testbed; proposal/intake/consensus flow; verified pack seeding; red-team privacy tests.",
    "Open specifications, local keys, transparent validator governance and independent build capability.":
        "Open specifications, local keys, independent Lightbearer Validators, transparent governance and reproducible builds.",
    "Current SAGE uses at least two consensus stages: CometBFT block commitment and a separate ≥2/3 weighted memory-vote quorum before memory status becomes committed [10–12].":
        "SAGE v11.18.4 uses at least two consensus stages: CometBFT block commitment and a separate ≥2/3 weighted memory-vote quorum before memory status becomes committed [10–12].",
    "Shows whether the object stays on the personal chain or is queued for public-chain submission.":
        "Shows whether material stays private or enters the explicit PUBLIC-CANDIDATE review and proposal flow.",
    "Carry compact public-chain announcements, research requests and control objects within strict payload and airtime budgets.":
        "Carry bounded availability hints, research messages and control objects within strict payload and airtime budgets.",
    "Private domains are encrypted at rest and never mounted by the Commons-chain transport.":
        "Private domains are encrypted at rest and never mounted by the Commons adapter, cache or transport services.",
    "Build handbook, reproducible image, disconnected-chain profile, curriculum and grant report.":
        "Build handbook, reproducible image, disconnected distribution profile, curriculum and grant report.",
}


ARCH = {
    **COMMON,
    "Sovereign personal chains with a separate governed public Commons chain":
        "Sovereign local AI with bounded access to a governed public Commons network",
    "Each SAGE Lantern is a sovereign local AI appliance. Its personal SAGE chain remains device-local; its Commons role retrieves from a separately stored cache of committed public knowledge. Connectivity improves freshness and collective knowledge but is not required for basic operation.":
        "Each SAGE Lantern is a sovereign local AI appliance. Personal SAGE remains device-local; an isolated Commons adapter provides proposer, query and replaceable cache functions only. The Lantern never runs a public validator or holds public consensus keys. Connectivity improves freshness but is not required for basic operation.",
    "Hard isolation: personal objects never appear in Commons inventories, transaction ranges or snapshots.":
        "Hard isolation: personal objects never appear in Commons proposals, queries, caches, packs or operational inventories.",
    "Resilient authority: Knowledge Point validators/full replicas participate in one public chain over normal internet where available.":
        "Resilient authority: stable independently operated Lightbearer Validators participate in public consensus; the Foundation must not control the validator population.",
    "PERSONAL SAGE INSTANCE / DATA ROOT\n└── Private domains\n│   ├── private.profile\n│   ├── private.conversations\n│   ├── private.health\n│   ├── private.preferences\n│   └── private.device-context\n\nCOMMONS CLIENT / ISOLATED PUBLIC CACHE\n└── Public domains\n│   ├── public.agriculture\n│   ├── public.medical\n│   ├── public.education\n│   ├── public.local-services\n│   └── public.community\n\nSEPARATE: quarantine + transient working state":
        "PERSONAL SAGE / PRIVATE DATA ROOT\n└── Private domains\n│   ├── private.profile\n│   ├── private.conversations\n│   ├── private.health\n│   ├── private.preferences\n│   └── private.device-context\n\nCOMMONS ADAPTER / ISOLATED CACHE\n├── PUBLIC-CANDIDATE outbox\n├── authenticated query client\n└── replaceable cached views\n    ├── public.agriculture\n    ├── public.medical\n    ├── public.education\n    └── public.local-services\n\nSEPARATE: quarantine + transient work",
    "Submit the returning candidate to the Commons chain and wait for committed status.":
        "If the user separately consents, create a PUBLIC-CANDIDATE, pass intake and make one fresh Lantern-authored submit.",
    "Private local context\n  HIV status + medicines + identity\n          │ local-only policy transform\n          ▼\nPublic request\n  domain: public.medical\n  query: contraindications for disease-x treatment\n  constraints: immunocompromised; interaction class ABC\n          │ mesh / Knowledge Point\n          ▼\nCandidate public knowledge\n          │ Commons-chain consensus\n          ▼\nLocal combination with private context":
        "Private local context\n  HIV status + medicines + identity\n          │ local-only minimisation\n          ▼\nPublic query\n  domain: public.medical\n  query: contraindications for disease-x treatment\n  constraints: broad, non-identifying classes\n          │ mesh / Lightbearer query service\n          ▼\nSourced public result / cached view\n          │ never auto-published\n          ▼\nLocal combination with private context",
    "Hyperlocal devices\nPersonal chain A     Personal chain B     Personal chain C\n      └──── isolated Commons caches + submission queues ────┐\n                                                            ▼\nKnowledge Point validators/full replicas ← internet → peer Knowledge Points\n                    ONE PUBLIC COMMONS CHAIN\n\nOptional: independent community chain ← SAGE Read/Copy federation → another chain":
        "SAGE Lanterns\nPersonal SAGE A      Personal SAGE B      Personal SAGE C\n      └──── proposer/query adapters + isolated caches ──────┐\n                                                            ▼\nCommons intake / Lightbearer services → independent validators\n                      PUBLIC SAGE COMMONS\n\nKnowledge Points may host isolated Lightbearer profiles; Lamplighters seed verified packs",
    "A SAGE Lantern may prepare a signed public.agriculture, public.medical or public.education submission while offline. It becomes a SAGE transaction only when broadcast to the Commons chain. CometBFT first commits the containing block; the memory then remains proposed until the SAGE validator-vote quorum commits or deprecates it. Knowledge Point validator membership is governed on-chain, not granted merely by being a physical Knowledge Point [11–13].":
        "A Lantern may create a signed logical PUBLIC-CANDIDATE envelope while offline. Intake verifies, deduplicates and quarantines it. After approval, the Lantern reconnects and creates exactly one fresh nonce/timestamp-bound SAGE submission; ambiguous timeouts are reconciled, never blindly resubmitted. Independently governed Lightbearer Validators determine canonical outcome [11–13].",
    "Do not label the envelope native SAGE federation. Native federation is a separate independent-chain Read/Copy protocol. The Commons transport carries an immutable signed candidate submission or an authenticated Commons-chain artifact and hands it to existing SAGE APIs after contact.":
        "The logical proposal envelope is Commons application data, not native SAGE federation and not a reusable signed HTTP request. It carries a proposer-authored candidate and provenance to durable intake. Any approved canonical submit uses a new SAGE authorization with a fresh nonce and timestamp after contact.",
    '{\n  "object_type": "candidate_submit | research | challenge | checkpoint",\n  "domain": "public.agriculture",\n  "object_id": "content-hash",\n  "author": "node-or-publisher-key",\n  "created_at": "2026-07-28T10:00:00Z",\n  "payload_ref": "inline-or-cid",\n  "provenance": ["source-id", "observation-id"],\n  "policy_version": "policy-hash",\n  "supersedes": [],\n  "contradicts": [],\n  "expiry": null,\n  "signature": "..."\n}':
        '{\n  "object_type": "PUBLIC-CANDIDATE",\n  "proposal_id": "uuid",\n  "domain": "public.agriculture",\n  "author": "lantern-proposer-key",\n  "created_at": "2026-08-11T10:00:00Z",\n  "public_derivative": "sanitized fact",\n  "provenance": ["public-source-id"],\n  "licence": "declared licence",\n  "policy_version": "policy-hash",\n  "signature": "logical-envelope-signature"\n}',
    "Private domains must not be representable in the public transport profile unless a distinct export protocol is invoked.":
        "Private sources must not be representable in the proposal schema; publication creates a separate sanitized derivative after explicit consent.",
    "The author signature binds content, domain, provenance and expiry; the Knowledge Point rechecks identity and SAGE RBAC before broadcast.":
        "The envelope signature binds candidate content, domain and provenance. Intake verifies proposer identity and policy; it is never reused as SAGE HTTP authorization.",
    "SAGE's existing BFT process remains the authority for public knowledge. Handhelds can carry candidate submissions while partitioned, but they must not pre-assign chain parent, block height or validator signatures. Only the active Commons chain orders the transaction. After inclusion, the SAGE memory-vote quorum determines committed or deprecated status [11,13].":
        "Public consensus remains the authority for canonical knowledge. Lanterns and Lamplighters may carry proposal envelopes while partitioned, but never assign height, validator signatures or outcome. Only stable Lightbearer Validators order and vote on canonical state; edge custody, intake approval and submission are distinct receipts [11,13].",
    "Maintain the personal chain in a separate process/data root from the Commons client/cache.":
        "Keep personal SAGE storage and authority isolated from the Commons adapter, outbox and replaceable cache.",
    "Knowledge Point validators/full replicas exchange transactions and blocks over normal internet and maintain redundant public-chain state.":
        "Independently operated Lightbearer Validators exchange transactions and blocks over stable links and maintain redundant canonical state.",
    "LoRa does not need to understand memory semantics. It enforces size, priority, airtime, retry and custody policy. Wi-Fi transports authenticated ranges and chunks with resume, deduplication and multi-peer completion.":
        "LoRa carries bounded discovery and control hints under size, priority and airtime policy. Wi-Fi, Ethernet and USB move verified content-addressed cache chunks and signed packs with resume, deduplication and multi-peer completion.",
    "Periodic beacons include scoped node ID, capability bits, permitted public-domain digests, highest cached Commons checkpoint and Wi-Fi rendezvous capability.":
        "Periodic beacons include scoped node ID, capability bits, permitted public-cache or pack summaries and bulk-transfer rendezvous capability.",
    "Peers request missing committed Commons ranges, snapshots and evidence by hash; multiple peers may supply different chunks.":
        "Peers request missing verified cache or pack chunks by hash; multiple peers may supply different chunks.",
    "Every range or snapshot is verified against the Commons chain commitment before becoming shareable.":
        "Every cache/pack artifact is verified against signed metadata, hashes, sequence and revocation state before activation or seeding.",
    "11. Knowledge Point architecture": "11. Knowledge Point and Lightbearer architecture",
    "Route to local-only, local-plus-SAGE, mesh research or queue-for-Knowledge-Point.":
        "Route to local-only, local-plus-cache, mesh research or queue-for-Lightbearer-service.",
    "Run two isolated logical SAGE roles: a device-local personal chain for private memory, and an intermittent client/cache for one public Commons SAGE chain whose validators and full replicas run at Knowledge Points. LoRa coordinates; Wi-Fi carries catch-up state, evidence and bulk bundles.":
        "Run personal SAGE locally and an isolated Commons proposer/query/cache adapter—not a second public SAGE process. Independent Lightbearer Validators govern public state. Lightbearer services answer or export; Lamplighters and peers distribute verified packs over bulk links while LoRa coordinates.",
    "SAGE already provides BFT memory consensus, sovereign independent-chain federation, same-chain peer join and same-chain state sync. Concurrent personal/public roles and delay-tolerant LoRa/Wi-Fi carriage are SAGE Commons integration work and must be proven in the pilot.":
        "SAGE v11.18.4 provides authenticated submission, consensus and defined federation/state-sync primitives. Commons proposal intake, external query, isolated cache, disconnected distribution and deterministic frozen-height committed-PUBLIC export are separately owned or required-SAGE-v12 contracts and must be proven.",
    "Submit eligible public candidates and ingest only verified committed Commons state.":
        "Queue signed PUBLIC-CANDIDATE envelopes, query authorized services and ingest only verified replaceable cache data.",
    "private, public, quarantine or transient.":
        "PRIVATE, PUBLIC-CANDIDATE, cache/public, quarantine or transient, with explicit SAGE numeric mapping.",
    "Checkpoint": "Cache/pack summary",
    "Commons chain head, validator set and compact state commitment.":
        "Bounded signed availability, sequence and checkpoint hints; never canonical proof by themselves.",
    "Committed catch-up range": "Verified cache chunk",
    "Missing public events between Commons checkpoints.":
        "Content-addressed derived public data from an authorized export service.",
    "Snapshot": "Signed pack",
    "Authorised same-chain bootstrap or repair.":
        "Knowledge, model, runtime or cache pack with signed metadata.",
    "Commons announcements and compact candidate submissions":
        "Commons availability hints and bounded research/control messages",
    "Committed Commons catch-up ranges and authorised snapshots":
        "Verified cache chunks, evidence and signed packs",
    "Commons-chain heads and checkpoint hashes":
        "Cache/pack sequence and checkpoint hints",
    "Knowledge Point validator/full-replica traffic":
        "Lamplighter and Lightbearer Cache/Seed transfers",
    "Research gateway": "Lightbearer Query/Export Service",
    "Accepts minimised public requests, performs constrained internet research and returns sourced candidate memories.":
        "Answers authenticated public queries or produces governed derived exports from an authorized full node; holds no validator key.",
    "Commons full replica": "Authorized public full node",
    "Maintains a full copy of the public Commons chain and serves authorised catch-up state.":
        "Maintains canonical public state for its exact service role; never exposes private or off-consensus data.",
    "Internet consensus peer": "Lightbearer Validator peer",
    "Replicates transactions and blocks with other Knowledge Points for resilience.":
        "Replicates transactions and blocks with independently operated validators across failure domains.",
    "Bundle cache": "Lightbearer Cache/Seed",
    "Stores models, software, Kiwix/ZIM, maps and evidence attachments.":
        "Stores and serves verified signed packs; custody grants no publishing or validation authority.",
    "Validator": "Lightbearer Validator",
    "Participates in the Commons CometBFT validator set and casts SAGE memory votes under on-chain governance.":
        "Independently operated, separately admitted and qualified; holds consensus keys under network governance.",
    "Mobile courier": "SAGE Lamplighter",
    "A travelling doctor, teacher or technician may carry a current public replica/cache and seed village peers; validator authority remains credential-bound.":
        "Untrusted offline super-seed carries verified signed packs and seeds peers over Wi-Fi, Ethernet or USB; never Root, publisher or validator.",
    "Local host": "Knowledge Point facility",
    "Community institution reduces recurring bandwidth, travel and support costs.":
        "Physical or institutional station may host isolated Lightbearer profiles; facility ownership grants no consensus power.",
    "Knowledge Points receive public questions and public objects. They do not receive a device's private retrieval context, identity graph or personal memory history.":
        "Knowledge Points, Lightbearers and Lamplighters receive only authorized public objects and minimized queries. They never receive private retrieval context, identity graphs or personal memory history.",
    "Authenticated HTTP/QUIC adapter with no personal-store interface.":
        "Authenticated, versioned adapter with no personal-store interface; proposal envelopes are not reusable HTTP requests.",
    "Clone base image; generate personal-chain identity/keys plus a separate Commons transport identity and storage root.":
        "Clone the base image; generate personal SAGE, proposer and transport identities on-device with isolated storage and no golden-image keys.",
    "Named stewards, transparent validator/publisher lists, incident procedure and appeal path.":
        "Named stewards, independent validator operators, transparent lists, incidents, separation of authority and appeals.",
    "Candidate submissions and committed Commons catch-up resume after interruption.":
        "Proposal delivery and verified cache/pack transfer resume after interruption without blind submission replay.",
    "Nodes exchange public candidate submissions and Commons-checkpoint summaries without private metadata.":
        "Nodes exchange bounded availability/control hints without private metadata; bulk proposals and packs use suitable links.",
    "Knowledge Point": "Knowledge Point / Lightbearer",
    "Research request, candidate submission, block consensus, memory-vote quorum and village catch-up complete end to end.":
        "Minimized query, proposal intake, fresh submit, independent validation and verified pack/cache delivery complete end to end.",
    "Opaque transport: LoRa and Wi-Fi move signed submissions, chain status and bundles but do not make consensus decisions.":
        "Opaque transport: LoRa carries bounded hints; suitable bulk links move proposals and verified packs. No transport makes consensus decisions.",
    "Independent heights: personal-chain height is unrelated to cached Commons-chain height.":
        "Independent state: personal-chain height is unrelated to any Commons cache checkpoint or pack sequence.",
    "Node/domain capability, chain head and inventory digest.":
        "Scoped capability plus bounded cache/pack availability summary and rendezvous hint.",
    "Queued public submission": "Logical proposal envelope",
    "Canonical statement, provenance, domain and author signature.":
        "Signed PUBLIC-CANDIDATE derivative and public provenance for durable intake.",
    "LoRa if compact; Wi-Fi otherwise": "Wi-Fi, Ethernet or USB; LoRa advertises availability only",
    "Queue candidate submissions and custody receipts during partitions; do not collect synthetic validator votes off-chain.":
        "Queue logical proposal envelopes and custody receipts during partitions; do not collect synthetic validator votes off-chain.",
    "Treat cached Commons height as sync state, not an independent authoritative chain.":
        "Treat a cache checkpoint or pack sequence as replaceable derived state, never an independent authoritative chain.",
    "One peer becomes a temporary access point based on power, storage, cached public height and policy.":
        "One peer becomes a temporary access point based on power, storage, verified cache/pack sequence and policy.",
    "Internet consensus peer": "Foundation Lightbearer",
    "Replicates transactions and blocks with other Knowledge Points for resilience.":
        "Operates disclosed public query, export, catalogue, intake or archive services; Foundation status grants no validator vote.",
    "Publisher agent": "Community Lightbearer",
    "Submits signed candidate memories from approved knowledge packs; it cannot unilaterally commit them.":
        "Runs stable community or institutional services under disclosed policy, credentials and exact authority boundaries.",
    "May traverse LoRa when compact": "Only bounded references/control hints traverse LoRa",
    "Normally Wi-Fi only": "Bulk payload uses Wi-Fi, Ethernet, USB or storage",
    "Search personal SAGE memory and committed Commons knowledge as distinct retrieval classes.":
        "Search personal SAGE and verified cached Commons views as distinct retrieval classes.",
    "Identify whether an answer uses private context, public memory, approved packs or remote research without exposing private details.":
        "Identify whether an answer uses private context, cached public knowledge, approved packs or remote research without exposing private details.",
    "Personal backup is user-controlled and separate from Commons snapshots/caches.":
        "Personal backup is user-controlled and separate from replaceable Commons caches and signed packs.",
    "SAGE tests: candidate broadcast, block inclusion, memory-vote quorum, challenge handling, lagging-cache catch-up and revocation propagation.":
        "SAGE tests: fresh submit, ambiguous-timeout reconciliation, block inclusion, memory-vote quorum, challenges and revocation propagation.",
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("kind", choices=("prd", "architecture"))
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    doc = Document(args.source)
    replace_exact(doc, PRD if args.kind == "prd" else ARCH)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
