#!/usr/bin/env python3
"""Print paragraph and table-cell text from a DOCX with stable indices."""

import sys
from docx import Document


def main(path: str) -> None:
    doc = Document(path)
    print(f"DOCUMENT {path}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"P{index:04d}\t{paragraph.style.name}\t{text}")
    for table_index, table in enumerate(doc.tables):
        print(f"TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
        for row_index, row in enumerate(table.rows):
            cells = [" | ".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
            print(f"T{table_index:02d}R{row_index:03d}\t" + " || ".join(cells))


if __name__ == "__main__":
    main(sys.argv[1])
